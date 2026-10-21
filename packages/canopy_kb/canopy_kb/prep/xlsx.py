"""Render a workbook so every number is readable next to what it means.

Each sheet becomes a markdown section. Each row with content becomes one line:

  - Row 11 · Percentage of total costs covered by the grantee (%, input): Burkina Faso 89%;
    Chad 90%; … · Source: GW analysis · Formula: = Inputs › Percentage of total costs …
    > Note (Chad): We assume that Malaria Consortium's projected budget in 2026 …

Column headers are read from the text rows above each column's first number (stacked
headers join: "Malaria Consortium / Nigeria / Kebbi"). Values use the cell's own number
format (0.8944 shown as 89%). Formulas are shown once per row, with cell references
replaced by the label of the row they point at. Cell comments (GiveWell's justifications)
and hyperlinks are kept.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

REF = re.compile(r"(?:(?:'([^']+)'|([A-Za-z0-9_][A-Za-z0-9_ .]*))!)?\$?([A-Z]{1,3})\$?(\d{1,7})(?![\d(])")
MAX_HEADER_ROWS = 12
#: A sheet past either size is a dataset (enrollment records, survey rows), not a model.
DATA_ROWS = 300
DATA_CHARS = 300_000


def fmt(v, nf: str) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)):
        nf = nf or "General"
        try:
            if "%" in nf:
                dec = len(nf.split(".")[1].split("%")[0]) if "." in nf else 0
                return f"{v * 100:.{min(dec, 4)}f}%"
            cur = "$" if "$" in nf else ""
            if nf != "General" and re.search(r"0(\.0+)?", nf):
                m = re.search(r"0\.(0+)", nf)
                dec = len(m.group(1)) if m else 0
                s = f"{v:,.{dec}f}" if "," in nf else f"{v:.{dec}f}"
                return f"{cur}{s}"
            if isinstance(v, float):
                if v != 0 and (abs(v) < 1e-3 or abs(v) >= 1e9):
                    return f"{cur}{v:.4g}"
                return f"{cur}{round(v, 6):,}" if abs(v) >= 1e4 else f"{cur}{round(v, 6)}"
            return f"{cur}{v:,}" if abs(v) >= 1e4 else f"{cur}{v}"
        except (ValueError, IndexError):
            return str(v)
    s = str(v).strip()
    return re.sub(r"\s+", " ", s)


def _is_text(v) -> bool:
    return isinstance(v, str) and v.strip() not in ("", "-") and not v.startswith("=")


def _headers(ws, rows) -> tuple[dict[int, str], set[int]]:
    """Column headers. The header band is the text rows just above where the sheet's data
    starts (the first row of any column holding three or more numbers); a numeric column
    with nothing in that band falls back to the text just above its own first number.
    Stacked headers join ("Nigeria / Kebbi"); a spanning parent label fills the blank cells
    to its right; a level shared by nearly every column ("Malaria Consortium") is dropped."""
    first_num: dict[int, int] = {}
    count: dict[int, int] = {}
    text_at: dict[int, list] = {}
    for row in rows:
        for c in row:
            v = c.value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                first_num.setdefault(c.column, c.row)
                count[c.column] = count.get(c.column, 0) + 1
            elif _is_text(v) and c.row <= 80:
                text_at.setdefault(c.column, []).append((c.row, v.strip()[:60]))
    data_cols = [c for c, n in count.items() if n >= 3]
    if not data_cols:
        return {}, set()
    r0 = min(first_num[c] for c in data_cols)
    band = sorted({r for col in text_at for r, _ in text_at[col] if r0 - MAX_HEADER_ROWS <= r < r0})
    levels: dict[int, dict[int, str]] = {}
    for r in band:
        levels[r] = {col: t for col, items in text_at.items() for rr, t in items if rr == r}
    cols = sorted(set(first_num) | {c for lv in levels.values() for c in lv})
    header: dict[int, list[str]] = {c: [] for c in cols}
    for r in band:
        lv = levels[r]
        last = ""
        for col in cols:
            if col in lv:
                last = lv[col]
                header[col].append(last)
            elif last and col in first_num and r != band[-1]:
                header[col].append(last)  # under a spanning label
            else:
                last = "" if col not in first_num else last
    for col in first_num:
        if not header.get(col):
            above = [t for r, t in text_at.get(col, []) if first_num[col] - MAX_HEADER_ROWS <= r < first_num[col]]
            if above:
                header[col] = above[-2:]
    headed = [c for c in header if header[c]]
    freq: dict[str, int] = {}
    for c in headed:
        for p in set(header[c]):
            freq[p] = freq.get(p, 0) + 1
    common = {p for p, n in freq.items() if len(headed) >= 5 and n >= 0.8 * len(headed)}
    out = {}
    for c in headed:
        parts = [p for p in dict.fromkeys(header[c]) if p not in common] or header[c][-1:]
        out[c] = " / ".join(parts)
    return out, set(band)


def render(path: Path, max_cell: int = 2000) -> tuple[str, list[str], list[str], str]:
    """Returns (markdown text, sheet names, hyperlink targets, dataset text)."""
    import io
    raw = Path(path).read_bytes()  # cached files carry no .xlsx extension, which openpyxl insists on for paths
    wf = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    wv = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    sheets = [ws.title for ws in wv.worksheets]
    # row labels for formula translation: sheet -> row -> label
    labels: dict[str, dict[int, str]] = {}
    grid: dict[str, list] = {}
    for ws in wv.worksheets:
        rows = list(ws.iter_rows())
        grid[ws.title] = rows
        lab = {}
        for row in rows:
            for c in row[:12]:
                if _is_text(c.value) and len(c.value.strip()) > 2:
                    lab[c.row] = c.value.strip()[:80]
                    break
        labels[ws.title] = lab
    links: list[str] = []
    out: list[str] = []
    data: list[str] = []
    for ws in wv.worksheets:
        fs = wf[ws.title]
        rows = grid[ws.title]
        if not rows:
            continue
        sheet_start = len(out)
        out.append(f"## Sheet: {ws.title}")
        header, header_rows = _headers(ws, rows)
        for row in rows:
            cells = [c for c in row if c.value not in (None, "")]
            if not cells:
                continue
            r = row[0].row
            label_cells, values, notes = [], [], []
            for c in cells:
                v = c.value
                link = getattr(getattr(c, "hyperlink", None), "target", None)
                if link:
                    links.append(link)
                if c.comment and c.comment.text.strip():
                    who = header.get(c.column) or c.coordinate
                    notes.append(f"  > Note ({who}): {re.sub(chr(10)+'+', ' ', c.comment.text.strip())[:max_cell]}")
                if isinstance(v, (int, float, dt.date)) and not isinstance(v, bool):
                    h = header.get(c.column, get_column_letter(c.column))
                    values.append(f"{h}: {fmt(v, c.number_format)}")
                elif _is_text(v) or isinstance(v, bool):
                    t = fmt(v, "")[:max_cell] + (f" <{link}>" if link else "")
                    if c.column in header and r > max(header_rows, default=0):
                        values.append(f"{header[c.column]}: {t}")
                    else:
                        label_cells.append(t)
            if not values and not label_cells and not notes:
                continue
            # one representative formula per row, references shown as the rows they name
            formula = ""
            for c in cells:
                fv = fs.cell(c.row, c.column).value
                if isinstance(fv, str) and fv.startswith("=") and isinstance(c.value, (int, float)):
                    def sub(m):
                        sh = m.group(1) or m.group(2) or ws.title
                        try:
                            rr = int(m.group(4))
                        except ValueError:
                            return m.group(0)
                        lab = labels.get(sh, {}).get(rr)
                        if not lab:
                            return m.group(0)
                        return f"[{sh} › {lab}]" if sh != ws.title else f"[{lab}]"
                    formula = REF.sub(sub, fv[1:])[:400]
                    break
            if r in header_rows and not values:
                line = f"- Row {r} (headers): " + " | ".join(label_cells)
            else:
                line = f"- Row {r} · " + " · ".join(label_cells) if label_cells else f"- Row {r}"
                if values:
                    line += ": " + "; ".join(values)
                if formula:
                    line += f" · Formula: ={formula}"
            out.append(line)
            out.extend(notes)
        body = out[sheet_start:]
        nrows = sum(1 for ln in body if ln.startswith("- Row"))
        if nrows > DATA_ROWS or sum(map(len, body)) > DATA_CHARS:
            # a record-level dataset, not model logic: full rows go to the data text,
            # the analysis text keeps a summary (what the columns are, how many rows, a sample)
            data.extend(body)
            hdr = [ln for ln in body if "(headers)" in ln][:3]
            sample = [ln[:600] for ln in body if ln.startswith("- Row") and "(headers)" not in ln][:5]
            del out[sheet_start:]
            out.append(f"## Sheet: {ws.title} (dataset: {nrows} rows; full rows in the workbook's data tabs)")
            out.extend(h[:1500] for h in hdr)
            out.append("Sample rows:")
            out.extend(sample)
    return "\n".join(out), sheets, sorted(set(links)), "\n".join(data)


if __name__ == "__main__":
    import sys
    text, sheets, links, data = render(Path(sys.argv[1]))
    print(text)
